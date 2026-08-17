from __future__ import annotations

import csv
import hashlib
import json
import math
import mimetypes
import os
import re
import shutil
import subprocess
import textwrap
import time
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator, Sequence

from PIL import Image, ImageDraw, ImageFile, ImageFont, ImageOps, UnidentifiedImageError

ImageFile.LOAD_TRUNCATED_IMAGES = True
Image.MAX_IMAGE_PIXELS = None

MODEL_EXTENSIONS = {
    ".glb", ".gltf", ".fbx", ".obj", ".blend", ".dae", ".stl", ".ply", ".3ds", ".abc",
    ".max", ".ma", ".mb", ".c4d", ".skp", ".x3d", ".lwo", ".vox", ".vrm", ".bvh",
    ".md5mesh", ".md5anim"
}
IMAGE_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".webp", ".tga", ".bmp", ".gif", ".tif", ".tiff",
    ".exr", ".hdr", ".dds", ".ktx", ".ktx2", ".svg", ".ico", ".psd", ".psb",
    ".ai", ".eps", ".xcf", ".kra", ".aseprite"
}
AUDIO_EXTENSIONS = {".wav", ".ogg", ".mp3", ".flac", ".m4a", ".aac", ".opus", ".wma", ".mid", ".midi"}
VIDEO_EXTENSIONS = {".mp4", ".webm", ".mov", ".avi", ".mkv", ".m4v"}
FONT_EXTENSIONS = {".ttf", ".otf", ".woff", ".woff2"}
PACKAGE_EXTENSIONS = {
    ".zip", ".7z", ".rar", ".tar", ".tgz", ".gz", ".pak", ".unitypackage", ".uasset",
    ".umap", ".asset", ".bundle", ".prefab", ".unity", ".mesh", ".skeleton", ".res"
}
MATERIAL_EXTENSIONS = {
    ".mtl", ".mat", ".material", ".shader", ".anim", ".controller", ".atlas", ".cubemap",
    ".tres", ".tscn", ".vmt", ".vmf"
}
ALL_ASSET_EXTENSIONS = (
    MODEL_EXTENSIONS | IMAGE_EXTENSIONS | AUDIO_EXTENSIONS | VIDEO_EXTENSIONS |
    FONT_EXTENSIONS | PACKAGE_EXTENSIONS | MATERIAL_EXTENSIONS
)

DEFAULT_EXCLUDED_DIRS = {
    ".git", ".github", ".idea", ".vscode", "node_modules", "dist", "build", "coverage",
    "__pycache__", ".pytest_cache", "03_Docs/Asset_Catalog", "Saved", "Intermediate", "Binaries"
}

TEXTURE_ROLE_PATTERNS: list[tuple[str, tuple[str, ...]]] = [
    ("base_color", ("albedo", "basecolor", "base_color", "diffuse", "_dif", "_diff", "color", "colour")),
    ("normal", ("normal", "_nrm", "_nor", "norm")),
    ("roughness", ("roughness", "rough", "_rgh")),
    ("metallic", ("metallic", "metalness", "metal", "_met")),
    ("ambient_occlusion", ("ambientocclusion", "ambient_occlusion", "_ao", "occlusion")),
    ("emission", ("emissive", "emission", "emit", "_emi")),
    ("opacity", ("opacity", "alpha", "transparency", "mask")),
    ("height", ("height", "displacement", "disp", "bump")),
    ("orm_packed", ("_orm", "occlusionroughnessmetallic", "arm")),
]

LICENSE_FILENAMES = {
    "license", "license.txt", "license.md", "licence", "licence.txt", "licence.md",
    "credits", "credits.txt", "credits.md", "third_party_notices.md", "readme.md", "readme.txt",
    "terms.txt", "terms.md", "eula.txt", "eula.md"
}


@dataclass(slots=True)
class SourceSpec:
    origin: str
    repository: str
    root: Path
    label: str
    include_roots: list[Path]


@dataclass(slots=True)
class CatalogConfig:
    output_dir: Path
    preview_size: int = 512
    image_assets_per_page: int = 2
    generic_assets_per_page: int = 2
    max_model_assets_per_volume: int = 80
    max_image_assets_per_volume: int = 240
    max_generic_assets_per_volume: int = 160
    include_preview_folder_in_zip: bool = False
    max_committable_file_mb: int = 92
    model_triangle_warning: int = 5000
    model_material_warning: int = 3
    texture_dimension_warning: int = 1024
    audio_bitrate_warning: int = 320_000


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def normalize_path(value: str | Path) -> str:
    return str(value).replace("\\", "/")


def slugify(value: str, max_length: int = 90) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    ascii_value = re.sub(r"[^A-Za-z0-9._-]+", "_", ascii_value).strip("._-")
    return (ascii_value or "asset")[:max_length]


def human_bytes(value: int | float | None) -> str:
    if value is None:
        return "N/D"
    number = float(value)
    units = ["B", "KB", "MB", "GB", "TB"]
    for unit in units:
        if abs(number) < 1024.0 or unit == units[-1]:
            return f"{number:.0f} {unit}" if unit == "B" else f"{number:.2f} {unit}"
        number /= 1024.0
    return f"{number:.2f} TB"


def human_duration(seconds: float | int | None) -> str:
    if seconds is None:
        return "N/D"
    try:
        total = float(seconds)
    except (TypeError, ValueError):
        return "N/D"
    minutes, sec = divmod(total, 60)
    hours, minutes = divmod(minutes, 60)
    if hours >= 1:
        return f"{int(hours):02d}:{int(minutes):02d}:{sec:05.2f}"
    return f"{int(minutes):02d}:{sec:05.2f}"


def sha256_file(path: Path, chunk_size: int = 4 * 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            chunk = handle.read(chunk_size)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def asset_kind(extension: str) -> str:
    ext = extension.lower()
    if ext in MODEL_EXTENSIONS:
        return "MODEL_3D"
    if ext in IMAGE_EXTENSIONS:
        return "IMAGE_TEXTURE"
    if ext in AUDIO_EXTENSIONS:
        return "AUDIO"
    if ext in VIDEO_EXTENSIONS:
        return "VIDEO"
    if ext in FONT_EXTENSIONS:
        return "FONT"
    if ext in PACKAGE_EXTENSIONS:
        return "PACKAGE_CONTAINER"
    if ext in MATERIAL_EXTENSIONS:
        return "MATERIAL_DATA"
    return "OTHER"


def infer_texture_role(path: Path) -> str:
    token = path.stem.lower().replace("-", "_").replace(" ", "_")
    compact = re.sub(r"[^a-z0-9_]", "", token)
    for role, patterns in TEXTURE_ROLE_PATTERNS:
        if any(pattern in compact for pattern in patterns):
            return role
    return "unspecified"


def infer_category(relative_path: Path, kind: str, origin: str, label: str) -> tuple[str, str, str]:
    parts = list(relative_path.parts)
    lower_parts = [part.lower() for part in parts]
    pack = label
    category = kind.title().replace("_", " ")
    subcategory = "Geral"

    if origin == "TEMPLATE":
        if parts and parts[0].lower() == "public":
            parts = parts[1:]
            lower_parts = lower_parts[1:]
        if parts:
            category = parts[0].replace("_", " ").title()
        if len(parts) > 1:
            subcategory = parts[1].replace("_", " ").title()
        pack = "World of Claudecraft"
    else:
        if "00_packs_raw" in lower_parts:
            index = lower_parts.index("00_packs_raw")
            if len(parts) > index + 1:
                pack = parts[index + 1]
            category = _category_from_tokens(lower_parts, kind)
            if len(parts) > index + 2:
                subcategory = parts[index + 2].replace("_", " ").title()
        elif "01_assets_organizados" in lower_parts:
            index = lower_parts.index("01_assets_organizados")
            if len(parts) > index + 1:
                category = parts[index + 1].replace("_", " ").title()
                pack = f"Assets Organizados - {category}"
            if len(parts) > index + 2:
                subcategory = parts[index + 2].replace("_", " ").title()
        elif "02_conceptart" in lower_parts:
            category = "Concept Art"
            subcategory = parts[-2].replace("_", " ").title() if len(parts) > 1 else "Geral"
            pack = "Concept Art Demons Whip"
        elif any("torment" in part for part in lower_parts):
            category = "Textures"
            subcategory = "Torment Textures"
            pack = "Torment Textures"
        else:
            category = _category_from_tokens(lower_parts, kind)

    return category, subcategory, pack


def _category_from_tokens(tokens: Sequence[str], kind: str) -> str:
    mapping = [
        ("character", "Characters"), ("player", "Characters"), ("hero", "Characters"),
        ("enemy", "Enemies"), ("enemies", "Enemies"), ("creature", "Enemies"),
        ("monster", "Enemies"), ("boss", "Bosses"), ("weapon", "Weapons"),
        ("sword", "Weapons"), ("axe", "Weapons"), ("shield", "Weapons"),
        ("armor", "Armor"), ("armour", "Armor"), ("building", "Buildings"),
        ("castle", "Buildings"), ("church", "Buildings"), ("house", "Buildings"),
        ("dungeon", "Dungeon"), ("grave", "Dungeon"), ("cemetery", "Dungeon"),
        ("nature", "Nature"), ("tree", "Nature"), ("foliage", "Nature"),
        ("rock", "Nature"), ("mount", "Mounts"), ("prop", "Props"),
        ("texture", "Textures"), ("material", "Textures"), ("ui", "UI"),
        ("icon", "UI"), ("vfx", "VFX"), ("effect", "VFX"), ("audio", "Audio"),
        ("music", "Audio"), ("sound", "Audio"),
    ]
    joined = "/".join(tokens)
    for token, category in mapping:
        if token in joined:
            return category
    return kind.title().replace("_", " ")


def should_skip(path: Path, source_root: Path) -> bool:
    try:
        relative = path.relative_to(source_root)
    except ValueError:
        relative = path
    normalized = normalize_path(relative).lower()
    for excluded in DEFAULT_EXCLUDED_DIRS:
        excluded_normalized = excluded.lower().replace("\\", "/")
        if normalized == excluded_normalized or normalized.startswith(excluded_normalized.rstrip("/") + "/"):
            return True
    if path.name.startswith(".~lock.") or path.name in {"Thumbs.db", ".DS_Store"}:
        return True
    return False


def discover_asset_files(spec: SourceSpec) -> Iterator[Path]:
    seen: set[Path] = set()
    for include_root in spec.include_roots:
        root = include_root if include_root.is_absolute() else spec.root / include_root
        if not root.exists():
            continue
        if root.is_file():
            candidates = [root]
        else:
            candidates = root.rglob("*")
        for path in candidates:
            if not path.is_file() or path in seen:
                continue
            if should_skip(path, spec.root):
                continue
            if path.suffix.lower() not in ALL_ASSET_EXTENSIONS:
                continue
            seen.add(path)
            yield path


def find_license_evidence(asset_path: Path, source_root: Path) -> tuple[str, str]:
    current = asset_path.parent
    for _ in range(24):
        if not current.exists():
            break
        try:
            files = {item.name.lower(): item for item in current.iterdir() if item.is_file()}
        except OSError:
            files = {}
        for candidate in LICENSE_FILENAMES:
            if candidate in files:
                evidence = normalize_path(files[candidate].relative_to(source_root))
                return "FOUND_NEARBY", evidence
        if current == source_root or source_root not in current.parents:
            break
        current = current.parent
    for project_candidate in [source_root / "LICENSE", source_root / "CREDITS.md", source_root / "THIRD_PARTY_NOTICES.md"]:
        if project_candidate.exists():
            return "PROJECT_LEVEL_ONLY", normalize_path(project_candidate.relative_to(source_root))
    return "MISSING_REVIEW_REQUIRED", ""


def build_base_records(specs: Sequence[SourceSpec]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    asset_number = 1
    for spec in specs:
        for path in discover_asset_files(spec):
            extension = path.suffix.lower()
            kind = asset_kind(extension)
            relative = path.relative_to(spec.root)
            category, subcategory, pack = infer_category(relative, kind, spec.origin, spec.label)
            try:
                size_bytes = path.stat().st_size
            except OSError:
                size_bytes = 0
            try:
                checksum = sha256_file(path)
            except OSError as exc:
                checksum = ""
            license_status, license_evidence = find_license_evidence(path, spec.root)
            record = {
                "asset_id": f"{spec.origin[:3]}-{asset_number:06d}",
                "origin": spec.origin,
                "repository": spec.repository,
                "source_label": spec.label,
                "source_root": normalize_path(spec.root),
                "relative_path": normalize_path(relative),
                "absolute_path": str(path.resolve()),
                "name": path.stem,
                "filename": path.name,
                "extension": extension,
                "mime_type": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
                "kind": kind,
                "category": category,
                "subcategory": subcategory,
                "pack": pack,
                "size_bytes": size_bytes,
                "sha256": checksum,
                "duplicate_count": 1,
                "duplicate_group": "",
                "is_duplicate": False,
                "technical_status": "PENDING",
                "preview_path": "",
                "license_status": license_status,
                "license_evidence": license_evidence,
                "analysis_notes": [],
                "warnings": [],
                "texture_role": infer_texture_role(path) if kind == "IMAGE_TEXTURE" else "",
                "generated_at": utc_now_iso(),
            }
            records.append(record)
            asset_number += 1

    by_hash: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        if record["sha256"]:
            by_hash[record["sha256"]].append(record)
    duplicate_counter = 1
    for checksum, group in by_hash.items():
        if len(group) < 2:
            continue
        group_id = f"DUP-{duplicate_counter:05d}"
        duplicate_counter += 1
        for record in group:
            record["duplicate_count"] = len(group)
            record["duplicate_group"] = group_id
            record["is_duplicate"] = True
    return records


def _font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = []
    if bold:
        candidates.extend([
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
        ])
    else:
        candidates.extend([
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        ])
    for candidate in candidates:
        if Path(candidate).exists():
            try:
                return ImageFont.truetype(candidate, size=size)
            except OSError:
                pass
    return ImageFont.load_default()


def make_placeholder_preview(output_path: Path, title: str, subtitle: str, size: int = 512) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (size, size), (31, 31, 36))
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((28, 28, size - 28, size - 28), radius=20, outline=(110, 110, 120), width=3)
    title_font = _font(max(20, size // 22), bold=True)
    body_font = _font(max(15, size // 31))
    wrapped_title = textwrap.wrap(title, width=24)[:4]
    y = size * 0.30
    for line in wrapped_title:
        bbox = draw.textbbox((0, 0), line, font=title_font)
        draw.text(((size - (bbox[2] - bbox[0])) / 2, y), line, fill=(235, 235, 238), font=title_font)
        y += bbox[3] - bbox[1] + 8
    y += 18
    for line in textwrap.wrap(subtitle, width=34)[:5]:
        bbox = draw.textbbox((0, 0), line, font=body_font)
        draw.text(((size - (bbox[2] - bbox[0])) / 2, y), line, fill=(175, 175, 184), font=body_font)
        y += bbox[3] - bbox[1] + 5
    image.save(output_path, "JPEG", quality=82, optimize=True)


def _checkerboard(width: int, height: int, tile: int = 24) -> Image.Image:
    background = Image.new("RGB", (width, height), (220, 220, 220))
    draw = ImageDraw.Draw(background)
    for y in range(0, height, tile):
        for x in range(0, width, tile):
            if ((x // tile) + (y // tile)) % 2:
                draw.rectangle((x, y, x + tile - 1, y + tile - 1), fill=(180, 180, 180))
    return background


def _safe_image_open(path: Path) -> Image.Image:
    if path.suffix.lower() == ".svg":
        import cairosvg
        png_bytes = cairosvg.svg2png(url=str(path))
        from io import BytesIO
        return Image.open(BytesIO(png_bytes))
    return Image.open(path)


def inspect_image(record: dict[str, Any], preview_dir: Path, preview_size: int) -> None:
    path = Path(record["absolute_path"])
    preview_path = preview_dir / f"{record['asset_id']}_{slugify(path.stem)}.jpg"
    try:
        with _safe_image_open(path) as original:
            image = original.copy()
            record["image_width"] = int(original.width)
            record["image_height"] = int(original.height)
            record["image_mode"] = original.mode
            record["image_format"] = original.format or path.suffix.lstrip(".").upper()
            record["image_frames"] = int(getattr(original, "n_frames", 1))
            bands = original.getbands()
            record["image_channels"] = len(bands)
            record["image_has_alpha"] = "A" in bands or "transparency" in original.info
            record["image_icc_profile"] = bool(original.info.get("icc_profile"))
            record["image_power_of_two"] = (
                original.width > 0 and original.height > 0 and
                (original.width & (original.width - 1) == 0) and
                (original.height & (original.height - 1) == 0)
            )
            bits_per_channel = 16 if "16" in original.mode or original.mode.startswith("I") else 8
            record["image_bit_depth_estimate"] = bits_per_channel
            record["decoded_memory_bytes_estimate"] = original.width * original.height * max(1, len(bands)) * (bits_per_channel // 8)

            thumb = image.convert("RGBA")
            thumb.thumbnail((preview_size - 48, preview_size - 48), Image.Resampling.LANCZOS)
            canvas = _checkerboard(preview_size, preview_size)
            x = (preview_size - thumb.width) // 2
            y = (preview_size - thumb.height) // 2
            canvas.paste(thumb.convert("RGB"), (x, y), thumb.getchannel("A"))
            canvas.save(preview_path, "JPEG", quality=84, optimize=True)

            sample = thumb.convert("RGB").resize((64, 64), Image.Resampling.BILINEAR)
            colors = sample.getcolors(maxcolors=64 * 64)
            record["sampled_color_count"] = len(colors) if colors else 4096
            if colors:
                dominant = sorted(colors, reverse=True)[:5]
                record["dominant_colors_rgb"] = [list(color) for _, color in dominant]
        record["preview_path"] = str(preview_path)
        record["technical_status"] = "OK"
    except Exception as exc:
        record["technical_status"] = "ERROR"
        record["analysis_notes"].append(f"Falha ao abrir/renderizar imagem: {type(exc).__name__}: {exc}")
        make_placeholder_preview(preview_path, path.name, "Imagem não pôde ser renderizada", preview_size)
        record["preview_path"] = str(preview_path)


def _run_ffprobe(path: Path) -> dict[str, Any]:
    command = [
        "ffprobe", "-v", "error", "-print_format", "json", "-show_format", "-show_streams", str(path)
    ]
    completed = subprocess.run(command, capture_output=True, text=True, timeout=120, check=False)
    if completed.returncode != 0:
        raise RuntimeError(completed.stderr.strip() or "ffprobe returned non-zero status")
    return json.loads(completed.stdout)


def inspect_audio(record: dict[str, Any], preview_dir: Path, preview_size: int) -> None:
    path = Path(record["absolute_path"])
    preview_path = preview_dir / f"{record['asset_id']}_{slugify(path.stem)}.png"
    try:
        probe = _run_ffprobe(path)
        audio_stream = next((stream for stream in probe.get("streams", []) if stream.get("codec_type") == "audio"), {})
        fmt = probe.get("format", {})
        record["audio_codec"] = audio_stream.get("codec_name", "")
        record["audio_channels"] = _int_or_none(audio_stream.get("channels"))
        record["audio_channel_layout"] = audio_stream.get("channel_layout", "")
        record["audio_sample_rate"] = _int_or_none(audio_stream.get("sample_rate"))
        record["audio_bit_rate"] = _int_or_none(audio_stream.get("bit_rate") or fmt.get("bit_rate"))
        record["audio_duration_seconds"] = _float_or_none(audio_stream.get("duration") or fmt.get("duration"))
        command = [
            "ffmpeg", "-y", "-v", "error", "-i", str(path),
            "-filter_complex", f"showwavespic=s={preview_size}x{preview_size}:colors=white",
            "-frames:v", "1", str(preview_path)
        ]
        completed = subprocess.run(command, capture_output=True, text=True, timeout=180, check=False)
        if completed.returncode != 0 or not preview_path.exists():
            make_placeholder_preview(preview_path.with_suffix(".jpg"), path.name, "Áudio - waveform indisponível", preview_size)
            preview_path = preview_path.with_suffix(".jpg")
        record["preview_path"] = str(preview_path)
        record["technical_status"] = "OK"
    except Exception as exc:
        record["technical_status"] = "ERROR"
        record["analysis_notes"].append(f"Falha na inspeção de áudio: {type(exc).__name__}: {exc}")
        fallback = preview_path.with_suffix(".jpg")
        make_placeholder_preview(fallback, path.name, "Áudio não pôde ser analisado", preview_size)
        record["preview_path"] = str(fallback)


def inspect_video(record: dict[str, Any], preview_dir: Path, preview_size: int) -> None:
    path = Path(record["absolute_path"])
    preview_path = preview_dir / f"{record['asset_id']}_{slugify(path.stem)}.jpg"
    try:
        probe = _run_ffprobe(path)
        video_stream = next((stream for stream in probe.get("streams", []) if stream.get("codec_type") == "video"), {})
        fmt = probe.get("format", {})
        duration = _float_or_none(video_stream.get("duration") or fmt.get("duration"))
        record["video_codec"] = video_stream.get("codec_name", "")
        record["video_width"] = _int_or_none(video_stream.get("width"))
        record["video_height"] = _int_or_none(video_stream.get("height"))
        record["video_duration_seconds"] = duration
        record["video_bit_rate"] = _int_or_none(video_stream.get("bit_rate") or fmt.get("bit_rate"))
        record["video_fps"] = _fraction_to_float(video_stream.get("avg_frame_rate") or video_stream.get("r_frame_rate"))
        seek = max(0.0, (duration or 0.0) * 0.25)
        command = [
            "ffmpeg", "-y", "-v", "error", "-ss", f"{seek:.3f}", "-i", str(path),
            "-frames:v", "1", "-vf", f"scale={preview_size}:{preview_size}:force_original_aspect_ratio=decrease,pad={preview_size}:{preview_size}:(ow-iw)/2:(oh-ih)/2",
            str(preview_path)
        ]
        completed = subprocess.run(command, capture_output=True, text=True, timeout=180, check=False)
        if completed.returncode != 0 or not preview_path.exists():
            make_placeholder_preview(preview_path, path.name, "Vídeo - frame indisponível", preview_size)
        record["preview_path"] = str(preview_path)
        record["technical_status"] = "OK"
    except Exception as exc:
        record["technical_status"] = "ERROR"
        record["analysis_notes"].append(f"Falha na inspeção de vídeo: {type(exc).__name__}: {exc}")
        make_placeholder_preview(preview_path, path.name, "Vídeo não pôde ser analisado", preview_size)
        record["preview_path"] = str(preview_path)


def inspect_font(record: dict[str, Any], preview_dir: Path, preview_size: int) -> None:
    path = Path(record["absolute_path"])
    preview_path = preview_dir / f"{record['asset_id']}_{slugify(path.stem)}.jpg"
    try:
        from fontTools.ttLib import TTFont
        font = TTFont(str(path), lazy=True)
        names = {}
        for name_record in font["name"].names:
            if name_record.nameID not in {1, 2, 4, 6}:
                continue
            try:
                value = name_record.toUnicode()
            except Exception:
                continue
            names.setdefault(name_record.nameID, value)
        record["font_family"] = names.get(1, "")
        record["font_subfamily"] = names.get(2, "")
        record["font_full_name"] = names.get(4, "")
        record["font_postscript_name"] = names.get(6, "")
        record["font_glyph_count"] = len(font.getGlyphOrder())
        font.close()

        canvas = Image.new("RGB", (preview_size, preview_size), (245, 242, 234))
        draw = ImageDraw.Draw(canvas)
        sample_font = ImageFont.truetype(str(path), size=max(24, preview_size // 10))
        small_font = ImageFont.truetype(str(path), size=max(16, preview_size // 18))
        sample_lines = ["Demons Whip", "Aa Bb Cc 012345", "Dark Fantasy PSX"]
        y = preview_size // 6
        for index, line in enumerate(sample_lines):
            font_obj = sample_font if index == 0 else small_font
            bbox = draw.textbbox((0, 0), line, font=font_obj)
            draw.text(((preview_size - (bbox[2] - bbox[0])) / 2, y), line, fill=(25, 25, 29), font=font_obj)
            y += (bbox[3] - bbox[1]) + preview_size // 12
        canvas.save(preview_path, "JPEG", quality=86, optimize=True)
        record["preview_path"] = str(preview_path)
        record["technical_status"] = "OK"
    except Exception as exc:
        record["technical_status"] = "PARTIAL"
        record["analysis_notes"].append(f"Fonte catalogada sem preview completo: {type(exc).__name__}: {exc}")
        make_placeholder_preview(preview_path, path.name, "Arquivo de fonte", preview_size)
        record["preview_path"] = str(preview_path)


def inspect_package_or_material(record: dict[str, Any], preview_dir: Path, preview_size: int) -> None:
    path = Path(record["absolute_path"])
    preview_path = preview_dir / f"{record['asset_id']}_{slugify(path.stem)}.jpg"
    kind = record["kind"]
    if path.suffix.lower() == ".zip":
        try:
            import zipfile
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
                record["container_member_count"] = len(members)
                record["container_uncompressed_bytes"] = sum(member.file_size for member in members)
                asset_members = [member.filename for member in members if Path(member.filename).suffix.lower() in ALL_ASSET_EXTENSIONS]
                record["container_asset_member_count"] = len(asset_members)
                record["container_sample_members"] = asset_members[:20]
                record["technical_status"] = "PARTIAL"
        except Exception as exc:
            record["analysis_notes"].append(f"ZIP não pôde ser enumerado: {type(exc).__name__}: {exc}")
            record["technical_status"] = "PARTIAL"
    elif kind == "MATERIAL_DATA" and record["size_bytes"] <= 2_000_000:
        try:
            sample = path.read_text(encoding="utf-8", errors="replace")[:4000]
            record["text_line_count_estimate"] = sample.count("\n") + 1
            record["text_preview"] = sample[:800]
            record["technical_status"] = "PARTIAL"
        except OSError as exc:
            record["technical_status"] = "ERROR"
            record["analysis_notes"].append(f"Falha ao ler material/dado: {exc}")
    else:
        record["technical_status"] = "UNSUPPORTED"
        record["analysis_notes"].append(
            "Formato-contêiner catalogado por metadados de arquivo. Geometria, rig e animações internos exigem fonte editável ou extrator específico."
        )
    subtitle = {
        "PACKAGE_CONTAINER": "Pacote/contêiner - inspeção interna limitada",
        "MATERIAL_DATA": "Material/dado auxiliar",
    }.get(kind, "Arquivo catalogado")
    make_placeholder_preview(preview_path, path.name, subtitle, preview_size)
    record["preview_path"] = str(preview_path)


def _int_or_none(value: Any) -> int | None:
    if value in (None, "", "N/A"):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _float_or_none(value: Any) -> float | None:
    if value in (None, "", "N/A"):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _fraction_to_float(value: Any) -> float | None:
    if value in (None, "", "0/0", "N/A"):
        return None
    try:
        if isinstance(value, str) and "/" in value:
            numerator, denominator = value.split("/", 1)
            denominator_value = float(denominator)
            return float(numerator) / denominator_value if denominator_value else None
        return float(value)
    except (TypeError, ValueError, ZeroDivisionError):
        return None


def inspect_non_model_assets(records: list[dict[str, Any]], preview_dir: Path, preview_size: int) -> None:
    preview_dir.mkdir(parents=True, exist_ok=True)
    for index, record in enumerate(records, start=1):
        kind = record["kind"]
        if kind == "MODEL_3D":
            continue
        if kind == "IMAGE_TEXTURE":
            inspect_image(record, preview_dir, preview_size)
        elif kind == "AUDIO":
            inspect_audio(record, preview_dir, preview_size)
        elif kind == "VIDEO":
            inspect_video(record, preview_dir, preview_size)
        elif kind == "FONT":
            inspect_font(record, preview_dir, preview_size)
        else:
            inspect_package_or_material(record, preview_dir, preview_size)


def build_model_task_manifest(records: Sequence[dict[str, Any]], preview_dir: Path) -> list[dict[str, Any]]:
    tasks: list[dict[str, Any]] = []
    representative_by_hash: dict[str, dict[str, Any]] = {}
    for record in records:
        if record["kind"] != "MODEL_3D":
            continue
        key = record.get("sha256") or record["absolute_path"]
        if key in representative_by_hash:
            representative = representative_by_hash[key]
            record["model_representative_asset_id"] = representative["asset_id"]
            continue
        representative_by_hash[key] = record
        preview_path = preview_dir / f"{record['asset_id']}_{slugify(Path(record['absolute_path']).stem)}.png"
        tasks.append({
            "asset_id": record["asset_id"],
            "path": record["absolute_path"],
            "preview_path": str(preview_path),
            "sha256": record.get("sha256", ""),
        })
    return tasks


def run_blender_model_inspection(
    records: list[dict[str, Any]],
    blender_script: Path,
    preview_dir: Path,
    work_dir: Path,
    blender_executable: str = "blender",
    chunk_size: int = 40,
    preview_size: int = 512,
) -> None:
    model_records = [record for record in records if record["kind"] == "MODEL_3D"]
    if not model_records:
        return
    preview_dir.mkdir(parents=True, exist_ok=True)
    work_dir.mkdir(parents=True, exist_ok=True)
    if shutil.which(blender_executable) is None:
        for record in model_records:
            preview_path = preview_dir / f"{record['asset_id']}_{slugify(Path(record['absolute_path']).stem)}.jpg"
            make_placeholder_preview(preview_path, record["filename"], "Blender não disponível - dados 3D pendentes", preview_size)
            record["preview_path"] = str(preview_path)
            record["technical_status"] = "UNSUPPORTED"
            record["analysis_notes"].append("Blender não estava disponível no ambiente de auditoria.")
        return

    tasks = build_model_task_manifest(model_records, preview_dir)
    metadata_by_asset: dict[str, dict[str, Any]] = {}
    for chunk_index, chunk in enumerate(_chunks(tasks, chunk_size), start=1):
        manifest_path = work_dir / f"blender_tasks_{chunk_index:04d}.json"
        result_path = work_dir / f"blender_results_{chunk_index:04d}.json"
        manifest_path.write_text(json.dumps(chunk, ensure_ascii=False, indent=2), encoding="utf-8")
        command = [
            blender_executable, "--background", "--factory-startup", "--python", str(blender_script), "--",
            "--manifest", str(manifest_path), "--output", str(result_path), "--preview-size", str(preview_size)
        ]
        completed = subprocess.run(command, capture_output=True, text=True, timeout=2400, check=False)
        log_path = work_dir / f"blender_chunk_{chunk_index:04d}.log"
        log_path.write_text(
            f"COMMAND: {' '.join(command)}\n\nSTDOUT:\n{completed.stdout}\n\nSTDERR:\n{completed.stderr}",
            encoding="utf-8", errors="replace"
        )
        if completed.returncode == 0 and result_path.exists():
            try:
                results = json.loads(result_path.read_text(encoding="utf-8"))
                for result in results:
                    metadata_by_asset[result["asset_id"]] = result
                continue
            except Exception:
                pass
        # Retry each asset individually to isolate corrupt/problematic source files.
        for task in chunk:
            single_manifest = work_dir / f"single_{task['asset_id']}.json"
            single_result = work_dir / f"single_{task['asset_id']}_result.json"
            single_manifest.write_text(json.dumps([task], ensure_ascii=False, indent=2), encoding="utf-8")
            single_command = [
                blender_executable, "--background", "--factory-startup", "--python", str(blender_script), "--",
                "--manifest", str(single_manifest), "--output", str(single_result), "--preview-size", str(preview_size)
            ]
            single_completed = subprocess.run(single_command, capture_output=True, text=True, timeout=420, check=False)
            if single_completed.returncode == 0 and single_result.exists():
                try:
                    result = json.loads(single_result.read_text(encoding="utf-8"))[0]
                    metadata_by_asset[result["asset_id"]] = result
                    continue
                except Exception:
                    pass
            metadata_by_asset[task["asset_id"]] = {
                "asset_id": task["asset_id"],
                "technical_status": "ERROR",
                "error": (single_completed.stderr or single_completed.stdout or "Blender inspection failed")[-2000:],
                "preview_path": task["preview_path"],
            }

    records_by_id = {record["asset_id"]: record for record in model_records}
    for representative_id, metadata in metadata_by_asset.items():
        record = records_by_id.get(representative_id)
        if record is None:
            continue
        merge_model_metadata(record, metadata, preview_size)

    representative_records = {record["asset_id"]: record for record in model_records}
    for record in model_records:
        representative_id = record.get("model_representative_asset_id")
        if not representative_id:
            continue
        representative = representative_records.get(representative_id)
        if representative is None:
            continue
        copy_keys = [
            key for key in representative.keys()
            if key.startswith("model_") or key in {
                "preview_path", "technical_status", "animation_names", "material_names", "texture_images",
                "analysis_notes", "warnings"
            }
        ]
        for key in copy_keys:
            value = representative.get(key)
            record[key] = json.loads(json.dumps(value, ensure_ascii=False)) if isinstance(value, (dict, list)) else value
        record["analysis_notes"].append(f"Metadados 3D reutilizados do duplicado idêntico {representative_id}.")


def merge_model_metadata(record: dict[str, Any], metadata: dict[str, Any], preview_size: int) -> None:
    for key, value in metadata.items():
        if key == "asset_id":
            continue
        record[key] = value
    preview_path = Path(str(record.get("preview_path") or metadata.get("preview_path") or ""))
    if not preview_path.exists():
        fallback = preview_path.with_suffix(".jpg") if str(preview_path) else Path(record["absolute_path"]).with_suffix(".preview.jpg")
        make_placeholder_preview(fallback, record["filename"], "Preview 3D indisponível", preview_size)
        record["preview_path"] = str(fallback)
    if metadata.get("error"):
        record.setdefault("analysis_notes", []).append(f"Falha na inspeção 3D: {metadata['error']}")


def _chunks(items: Sequence[Any], size: int) -> Iterator[list[Any]]:
    for index in range(0, len(items), max(1, size)):
        yield list(items[index:index + max(1, size)])


def apply_audit_warnings(records: list[dict[str, Any]], config: CatalogConfig) -> None:
    for record in records:
        warnings = record.setdefault("warnings", [])
        if record["kind"] == "MODEL_3D":
            triangles = _int_or_none(record.get("model_triangles"))
            materials = _int_or_none(record.get("model_material_count"))
            if triangles is not None and triangles > config.model_triangle_warning:
                warnings.append(f"Triângulos acima do alerta ({triangles:,} > {config.model_triangle_warning:,}).")
            if materials is not None and materials > config.model_material_warning:
                warnings.append(f"Materiais acima do alerta ({materials} > {config.model_material_warning}).")
            if record.get("model_animation_count", 0) and not record.get("model_has_rig"):
                warnings.append("Possui animações detectadas, mas nenhum rig/armature foi identificado.")
            if record.get("model_missing_texture_count", 0):
                warnings.append(f"Texturas ausentes detectadas: {record['model_missing_texture_count']}.")
            dimensions = record.get("model_dimensions") or []
            if dimensions and max([abs(float(value)) for value in dimensions if value is not None] or [0]) > 10_000:
                warnings.append("Dimensões extremamente altas; revisar unidade e escala.")
        elif record["kind"] == "IMAGE_TEXTURE":
            width = _int_or_none(record.get("image_width")) or 0
            height = _int_or_none(record.get("image_height")) or 0
            if max(width, height) > config.texture_dimension_warning:
                warnings.append(
                    f"Resolução acima do alerta ({width}x{height}; limite de referência {config.texture_dimension_warning}px)."
                )
            if record.get("texture_role") not in {"unspecified", "normal", "hdr"} and record.get("image_power_of_two") is False:
                warnings.append("Textura não é power-of-two.")
        elif record["kind"] == "AUDIO":
            bitrate = _int_or_none(record.get("audio_bit_rate"))
            if bitrate is not None and bitrate > config.audio_bitrate_warning:
                warnings.append(f"Bitrate acima do alerta ({bitrate:,} bps).")
        if record.get("license_status") == "MISSING_REVIEW_REQUIRED":
            warnings.append("Licença/fonte não localizada automaticamente; revisão legal necessária.")
        if record.get("is_duplicate"):
            warnings.append(f"Duplicado exato: grupo {record.get('duplicate_group')} ({record.get('duplicate_count')} cópias).")


def strip_runtime_paths(record: dict[str, Any], output_root: Path | None = None) -> dict[str, Any]:
    cleaned = dict(record)
    cleaned.pop("absolute_path", None)
    cleaned.pop("source_root", None)
    preview = cleaned.get("preview_path")
    if preview:
        preview_path = Path(str(preview))
        if output_root is not None:
            try:
                preview_path = preview_path.resolve().relative_to(output_root.resolve())
            except (ValueError, OSError):
                pass
        cleaned["preview_path"] = normalize_path(preview_path)
    return cleaned


def write_inventory_files(records: list[dict[str, Any]], data_dir: Path) -> dict[str, Path]:
    data_dir.mkdir(parents=True, exist_ok=True)
    json_path = data_dir / "asset_inventory_full.json"
    cleaned = [strip_runtime_paths(record, data_dir.parent) for record in records]
    json_path.write_text(json.dumps(cleaned, ensure_ascii=False, indent=2), encoding="utf-8")

    all_keys: list[str] = sorted({key for record in cleaned for key in record.keys()})
    preferred = [
        "asset_id", "origin", "repository", "source_label", "relative_path", "name", "filename", "extension",
        "kind", "category", "subcategory", "pack", "size_bytes", "sha256", "duplicate_count", "duplicate_group",
        "technical_status", "license_status", "license_evidence", "preview_path"
    ]
    fieldnames = preferred + [key for key in all_keys if key not in preferred]
    csv_path = data_dir / "asset_inventory_full.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for record in cleaned:
            row = {}
            for key in fieldnames:
                value = record.get(key, "")
                if isinstance(value, (list, dict)):
                    value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
                row[key] = value
            writer.writerow(row)

    duplicate_path = data_dir / "duplicate_groups.csv"
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in cleaned:
        if record.get("duplicate_group"):
            groups[record["duplicate_group"]].append(record)
    with duplicate_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["duplicate_group", "sha256", "copy_count", "asset_id", "origin", "relative_path", "size_bytes"])
        for group_id, group_records in sorted(groups.items()):
            for record in group_records:
                writer.writerow([
                    group_id, record.get("sha256", ""), len(group_records), record["asset_id"], record["origin"],
                    record["relative_path"], record["size_bytes"]
                ])

    replacement_path = data_dir / "replacement_matrix_template.csv"
    template_records = [record for record in cleaned if record["origin"] == "TEMPLATE"]
    with replacement_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([
            "template_asset_id", "template_path", "template_category", "template_kind", "candidate_new_asset_id",
            "candidate_new_path", "decision", "priority", "geometry_action", "texture_action", "rig_action",
            "animation_action", "scale_pivot_action", "license_gate", "runtime_evidence", "notes"
        ])
        for record in template_records:
            writer.writerow([
                record["asset_id"], record["relative_path"], record["category"], record["kind"], "", "", "PENDING",
                "", "", "", "", "", "", record.get("license_status", ""), "", ""
            ])

    unsupported_path = data_dir / "unsupported_or_partial_assets.csv"
    with unsupported_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["asset_id", "origin", "kind", "status", "relative_path", "notes"])
        for record in cleaned:
            if record.get("technical_status") in {"PARTIAL", "UNSUPPORTED", "ERROR"}:
                writer.writerow([
                    record["asset_id"], record["origin"], record["kind"], record.get("technical_status"),
                    record["relative_path"], " | ".join(record.get("analysis_notes", []))
                ])

    return {
        "json": json_path,
        "csv": csv_path,
        "duplicates": duplicate_path,
        "replacement_matrix": replacement_path,
        "unsupported": unsupported_path,
    }


def summarize_records(records: Sequence[dict[str, Any]]) -> dict[str, Any]:
    by_origin = Counter(record["origin"] for record in records)
    by_kind = Counter(record["kind"] for record in records)
    by_origin_kind = Counter((record["origin"], record["kind"]) for record in records)
    by_status = Counter(record.get("technical_status", "UNKNOWN") for record in records)
    by_category = Counter((record["origin"], record["category"]) for record in records)
    warnings = sum(len(record.get("warnings", [])) for record in records)
    duplicates = sum(1 for record in records if record.get("is_duplicate"))
    unique_hashes = len({record.get("sha256") for record in records if record.get("sha256")})
    total_size = sum(int(record.get("size_bytes", 0) or 0) for record in records)
    return {
        "generated_at": utc_now_iso(),
        "total_assets": len(records),
        "total_size_bytes": total_size,
        "by_origin": dict(sorted(by_origin.items())),
        "by_kind": dict(sorted(by_kind.items())),
        "by_origin_kind": [
            {"origin": origin, "kind": kind, "count": count}
            for (origin, kind), count in sorted(by_origin_kind.items())
        ],
        "by_status": dict(sorted(by_status.items())),
        "by_category": [
            {"origin": origin, "category": category, "count": count}
            for (origin, category), count in sorted(by_category.items())
        ],
        "warning_count": warnings,
        "duplicate_asset_count": duplicates,
        "unique_content_hash_count": unique_hashes,
    }


def write_summary_json(records: Sequence[dict[str, Any]], data_dir: Path) -> Path:
    summary = summarize_records(records)
    path = data_dir / "audit_summary.json"
    path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def write_readme(output_dir: Path, summary: dict[str, Any], volume_rows: list[dict[str, Any]]) -> Path:
    path = output_dir / "README.md"
    lines = [
        "# Demons Whip - Catálogo Técnico e Visual de Assets",
        "",
        f"Gerado em: `{summary['generated_at']}`",
        "",
        "Este pacote separa os assets do template **World of Claudecraft** dos assets novos do projeto **Demons Whip**.",
        "Os dados foram medidos diretamente nos arquivos disponíveis. Campos impossíveis de extrair automaticamente são marcados como `PARTIAL`, `UNSUPPORTED` ou `ERROR`.",
        "",
        "## Conteúdo",
        "",
        "- `Demons_Whip_Inventario_Mestre.docx/.pdf`: resumo executivo, metodologia, riscos e índice.",
        "- `volumes/`: catálogos visuais detalhados, com imagem e dados técnicos de cada asset.",
        "- `data/asset_inventory_full.csv/.json`: inventário estruturado completo.",
        "- `data/duplicate_groups.csv`: duplicações exatas detectadas por SHA-256.",
        "- `data/replacement_matrix_template.csv`: matriz pronta para a próxima etapa de substituição e normalização.",
        "- `data/unsupported_or_partial_assets.csv`: itens que exigem inspeção manual ou ferramenta específica.",
        "",
        "## Totais",
        "",
        f"- Assets catalogados: **{summary['total_assets']:,}**",
        f"- Tamanho somado: **{human_bytes(summary['total_size_bytes'])}**",
        f"- Conteúdos únicos por hash: **{summary['unique_content_hash_count']:,}**",
        f"- Cópias em grupos duplicados: **{summary['duplicate_asset_count']:,}**",
        f"- Alertas técnicos: **{summary['warning_count']:,}**",
        "",
        "## Volumes",
        "",
        "| Arquivo | Origem | Categoria | Tipo | Assets |",
        "|---|---|---|---|---:|",
    ]
    for row in volume_rows:
        lines.append(
            f"| `{row['docx_name']}` | {row['origin']} | {row['category']} | {row['kind']} | {row['asset_count']} |"
        )
    lines.extend([
        "",
        "## Limitações técnicas importantes",
        "",
        "- Arquivos compilados/contêineres (`.pak`, `.uasset`, `.unitypackage`, etc.) não revelam automaticamente polígonos, rig ou animações sem a fonte editável ou extrator específico.",
        "- O status de licença é uma detecção documental; aprovação jurídica continua necessária antes da publicação.",
        "- A detecção de root motion, LOD e collider é baseada em dados e convenções de nomes e deve ser confirmada no runtime.",
        "- A matriz de substituição é deliberadamente deixada sem decisões: a próxima etapa será feita sobre este inventário auditado.",
        "",
    ])
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


# ----------------------------- DOCX REPORTING -----------------------------

def _import_docx():
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor
    return {
        "Document": Document,
        "WD_SECTION": WD_SECTION,
        "WD_CELL_VERTICAL_ALIGNMENT": WD_CELL_VERTICAL_ALIGNMENT,
        "WD_TABLE_ALIGNMENT": WD_TABLE_ALIGNMENT,
        "WD_ALIGN_PARAGRAPH": WD_ALIGN_PARAGRAPH,
        "WD_BREAK": WD_BREAK,
        "OxmlElement": OxmlElement,
        "qn": qn,
        "Cm": Cm,
        "Inches": Inches,
        "Pt": Pt,
        "RGBColor": RGBColor,
    }


def _set_cell_shading(cell, fill: str) -> None:
    tools = _import_docx()
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(tools["qn"]("w:shd"))
    if shd is None:
        shd = tools["OxmlElement"]("w:shd")
        tc_pr.append(shd)
    shd.set(tools["qn"]("w:fill"), fill)


def _set_cell_margins(cell, top: int = 80, start: int = 80, bottom: int = 80, end: int = 80) -> None:
    tools = _import_docx()
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = tools["OxmlElement"]("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(tools["qn"](f"w:{margin}"))
        if node is None:
            node = tools["OxmlElement"](f"w:{margin}")
            tc_mar.append(node)
        node.set(tools["qn"]("w:w"), str(value))
        node.set(tools["qn"]("w:type"), "dxa")


def _set_repeat_table_header(row) -> None:
    tools = _import_docx()
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tools["OxmlElement"]("w:tblHeader")
    tbl_header.set(tools["qn"]("w:val"), "true")
    tr_pr.append(tbl_header)


def _add_page_number(paragraph) -> None:
    tools = _import_docx()
    run = paragraph.add_run()
    fld_char_1 = tools["OxmlElement"]("w:fldChar")
    fld_char_1.set(tools["qn"]("w:fldCharType"), "begin")
    instr_text = tools["OxmlElement"]("w:instrText")
    instr_text.set(tools["qn"]("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = tools["OxmlElement"]("w:fldChar")
    fld_char_2.set(tools["qn"]("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def configure_document(document, title: str, subtitle: str = "") -> None:
    tools = _import_docx()
    section = document.sections[0]
    section.top_margin = tools["Cm"](1.35)
    section.bottom_margin = tools["Cm"](1.25)
    section.left_margin = tools["Cm"](1.35)
    section.right_margin = tools["Cm"](1.35)
    section.header_distance = tools["Cm"](0.45)
    section.footer_distance = tools["Cm"](0.45)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = tools["Pt"](8.5)
    normal.paragraph_format.space_after = tools["Pt"](2.5)
    for style_name, size, color in [
        ("Title", 28, "251C1D"), ("Subtitle", 12, "6A4A4E"),
        ("Heading 1", 19, "4A171D"), ("Heading 2", 14, "6D232B"), ("Heading 3", 11, "7A3940")
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = tools["Pt"](size)
        style.font.color.rgb = tools["RGBColor"].from_string(color)

    header = section.header.paragraphs[0]
    header.alignment = tools["WD_ALIGN_PARAGRAPH"].CENTER
    run = header.add_run("PIXLAND  •  DEMONS WHIP  •  CATÁLOGO DE ASSETS")
    run.font.name = "Arial"
    run.font.size = tools["Pt"](8)
    run.font.color.rgb = tools["RGBColor"].from_string("6D6A70")

    footer = section.footer.paragraphs[0]
    footer.alignment = tools["WD_ALIGN_PARAGRAPH"].CENTER
    run = footer.add_run("Uso interno • ")
    run.font.name = "Arial"
    run.font.size = tools["Pt"](8)
    _add_page_number(footer)

    core = document.core_properties
    core.title = title
    core.subject = subtitle
    core.author = "Pixlland"
    core.keywords = "Demons Whip, asset catalog, technical audit, PSX, Pixlland"


def add_cover(document, title: str, subtitle: str, metadata_lines: Sequence[str]) -> None:
    tools = _import_docx()
    for _ in range(4):
        document.add_paragraph()
    paragraph = document.add_paragraph()
    paragraph.alignment = tools["WD_ALIGN_PARAGRAPH"].CENTER
    run = paragraph.add_run(title.upper())
    run.bold = True
    run.font.name = "Arial"
    run.font.size = tools["Pt"](30)
    run.font.color.rgb = tools["RGBColor"].from_string("44151A")

    paragraph = document.add_paragraph()
    paragraph.alignment = tools["WD_ALIGN_PARAGRAPH"].CENTER
    run = paragraph.add_run(subtitle)
    run.font.name = "Arial"
    run.font.size = tools["Pt"](13)
    run.font.color.rgb = tools["RGBColor"].from_string("6D4C50")

    document.add_paragraph()
    table = document.add_table(rows=len(metadata_lines), cols=1)
    table.alignment = tools["WD_TABLE_ALIGNMENT"].CENTER
    for row, line in zip(table.rows, metadata_lines):
        cell = row.cells[0]
        _set_cell_shading(cell, "F2ECE7")
        _set_cell_margins(cell, top=110, bottom=110, start=160, end=160)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = tools["WD_ALIGN_PARAGRAPH"].CENTER
        run = paragraph.add_run(line)
        run.font.name = "Arial"
        run.font.size = tools["Pt"](10)
    document.add_page_break()


def add_key_value_table(
    document,
    rows: Sequence[tuple[str, Any]],
    label_width_cm: float = 5.0,
    font_size: float = 8.0,
    cell_margin: int = 80,
) -> None:
    tools = _import_docx()
    table = document.add_table(rows=0, cols=2)
    table.alignment = tools["WD_TABLE_ALIGNMENT"].CENTER
    table.style = "Table Grid"
    for label, value in rows:
        row = table.add_row()
        row.cells[0].width = tools["Cm"](label_width_cm)
        _set_cell_shading(row.cells[0], "EDE2DC")
        _set_cell_margins(row.cells[0], top=cell_margin, start=cell_margin, bottom=cell_margin, end=cell_margin)
        _set_cell_margins(row.cells[1], top=cell_margin, start=cell_margin, bottom=cell_margin, end=cell_margin)
        label_run = row.cells[0].paragraphs[0].add_run(str(label))
        label_run.bold = True
        label_run.font.size = tools["Pt"](font_size)
        value_text = _value_to_text(value)
        value_run = row.cells[1].paragraphs[0].add_run(value_text)
        value_run.font.size = tools["Pt"](font_size)


def _value_to_text(value: Any) -> str:
    if value is None or value == "":
        return "N/D"
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    if isinstance(value, list):
        return ", ".join(str(item) for item in value) if value else "Nenhum"
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def asset_technical_rows(record: dict[str, Any]) -> list[tuple[str, Any]]:
    common = [
        ("ID", record.get("asset_id")),
        ("Origem", "Template" if record.get("origin") == "TEMPLATE" else "Asset novo"),
        ("Pack/Fonte", record.get("pack")),
        ("Categoria", f"{record.get('category')} / {record.get('subcategory')}"),
        ("Arquivo", record.get("filename")),
        ("Caminho", record.get("relative_path")),
        ("Formato", record.get("extension", "").lstrip(".").upper()),
        ("Tamanho", human_bytes(record.get("size_bytes"))),
        ("SHA-256", (record.get("sha256") or "")[:20] + ("…" if record.get("sha256") else "")),
        ("Status técnico", record.get("technical_status")),
        ("Licença", f"{record.get('license_status')} — {record.get('license_evidence') or 'sem evidência localizada'}"),
    ]
    kind = record.get("kind")
    if kind == "MODEL_3D":
        dimensions = record.get("model_dimensions") or []
        dimensions_text = " × ".join(f"{float(value):.4g}" for value in dimensions) if dimensions else "N/D"
        common.extend([
            ("Objetos / meshes", f"{record.get('model_object_count', 'N/D')} / {record.get('model_mesh_count', 'N/D')}"),
            ("Vértices", _format_int(record.get("model_vertices"))),
            ("Faces / triângulos", f"{_format_int(record.get('model_faces'))} / {_format_int(record.get('model_triangles'))}"),
            ("Materiais", f"{record.get('model_material_count', 'N/D')} — {', '.join(record.get('material_names', [])[:8])}"),
            ("UV sets / vertex colors", f"{record.get('model_uv_layer_count', 'N/D')} / {record.get('model_vertex_color_layer_count', 'N/D')}"),
            ("Dimensões XYZ", dimensions_text),
            ("Rig / armatures / bones", f"{_yes_no(record.get('model_has_rig'))} / {record.get('model_armature_count', 'N/D')} / {record.get('model_bone_count', 'N/D')}"),
            ("Skinned meshes", record.get("model_skinned_mesh_count", "N/D")),
            ("Animações", f"{record.get('model_animation_count', 'N/D')} — {', '.join(record.get('animation_names', [])[:12])}"),
            ("Duração total clips", human_duration(record.get("model_animation_total_seconds"))),
            ("Root motion detectado", _yes_no(record.get("model_root_motion_detected"))),
            ("LODs / colliders", f"{', '.join(record.get('model_lod_names', [])[:8]) or 'Nenhum'} / {', '.join(record.get('model_collider_names', [])[:8]) or 'Nenhum'}"),
            ("Texturas vinculadas", f"{record.get('model_texture_image_count', 'N/D')} ({record.get('model_missing_texture_count', 0)} ausentes)"),
        ])
    elif kind == "IMAGE_TEXTURE":
        common.extend([
            ("Papel da textura", record.get("texture_role")),
            ("Resolução", f"{record.get('image_width', 'N/D')} × {record.get('image_height', 'N/D')} px"),
            ("Modo / canais / alpha", f"{record.get('image_mode', 'N/D')} / {record.get('image_channels', 'N/D')} / {_yes_no(record.get('image_has_alpha'))}"),
            ("Bit depth estimado", record.get("image_bit_depth_estimate", "N/D")),
            ("Power-of-two", _yes_no(record.get("image_power_of_two"))),
            ("Memória decodificada estimada", human_bytes(record.get("decoded_memory_bytes_estimate"))),
            ("Frames", record.get("image_frames", "N/D")),
            ("Cores amostradas", record.get("sampled_color_count", "N/D")),
        ])
    elif kind == "AUDIO":
        common.extend([
            ("Codec", record.get("audio_codec")),
            ("Duração", human_duration(record.get("audio_duration_seconds"))),
            ("Canais / layout", f"{record.get('audio_channels', 'N/D')} / {record.get('audio_channel_layout', 'N/D')}"),
            ("Sample rate", f"{_format_int(record.get('audio_sample_rate'))} Hz"),
            ("Bitrate", f"{_format_int(record.get('audio_bit_rate'))} bps"),
        ])
    elif kind == "VIDEO":
        common.extend([
            ("Codec", record.get("video_codec")),
            ("Resolução", f"{record.get('video_width', 'N/D')} × {record.get('video_height', 'N/D')} px"),
            ("Duração", human_duration(record.get("video_duration_seconds"))),
            ("FPS", record.get("video_fps", "N/D")),
            ("Bitrate", f"{_format_int(record.get('video_bit_rate'))} bps"),
        ])
    elif kind == "FONT":
        common.extend([
            ("Família", record.get("font_family")),
            ("Subfamília", record.get("font_subfamily")),
            ("Nome completo", record.get("font_full_name")),
            ("Glifos", _format_int(record.get("font_glyph_count"))),
        ])
    else:
        common.extend([
            ("Membros no contêiner", record.get("container_member_count", "N/D")),
            ("Assets internos detectados", record.get("container_asset_member_count", "N/D")),
            ("Tamanho descompactado", human_bytes(record.get("container_uncompressed_bytes"))),
        ])
    if record.get("warnings"):
        common.append(("Alertas", " | ".join(record["warnings"])))
    if record.get("analysis_notes"):
        common.append(("Notas", " | ".join(record["analysis_notes"])))
    return common


def _format_int(value: Any) -> str:
    parsed = _int_or_none(value)
    return f"{parsed:,}" if parsed is not None else "N/D"


def _yes_no(value: Any) -> str:
    if value is None or value == "":
        return "N/D"
    return "Sim" if bool(value) else "Não"


def add_asset_full_page(document, record: dict[str, Any]) -> None:
    tools = _import_docx()
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = tools["Pt"](4)
    run = paragraph.add_run(f"{record['asset_id']}  •  {record['name']}")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = tools["Pt"](16)
    run.font.color.rgb = tools["RGBColor"].from_string("4A171D")
    status_run = paragraph.add_run(f"   [{record.get('technical_status', 'N/D')}]")
    status_run.font.size = tools["Pt"](9)
    status_run.font.color.rgb = tools["RGBColor"].from_string("7A6265")

    preview = Path(record.get("preview_path", ""))
    if preview.exists():
        picture_paragraph = document.add_paragraph()
        picture_paragraph.alignment = tools["WD_ALIGN_PARAGRAPH"].CENTER
        try:
            picture_paragraph.add_run().add_picture(str(preview), width=tools["Inches"](4.05))
        except Exception:
            pass
    add_key_value_table(
        document, asset_technical_rows(record), label_width_cm=5.0, font_size=7.0, cell_margin=38
    )
    document.add_page_break()


def add_asset_compact_card(cell, record: dict[str, Any]) -> None:
    tools = _import_docx()
    cell.vertical_alignment = tools["WD_CELL_VERTICAL_ALIGNMENT"].TOP
    _set_cell_margins(cell, top=70, start=75, bottom=70, end=75)
    title = cell.paragraphs[0]
    title.paragraph_format.space_after = tools["Pt"](2)
    run = title.add_run(f"{record['asset_id']} • {record['name']}")
    run.bold = True
    run.font.size = tools["Pt"](8.5)
    run.font.color.rgb = tools["RGBColor"].from_string("4A171D")
    preview = Path(record.get("preview_path", ""))
    if preview.exists():
        picture_paragraph = cell.add_paragraph()
        picture_paragraph.alignment = tools["WD_ALIGN_PARAGRAPH"].CENTER
        try:
            picture_paragraph.add_run().add_picture(str(preview), width=tools["Inches"](2.25))
        except Exception:
            pass
    compact_rows = asset_technical_rows(record)
    keep_labels = {
        "Origem", "Pack/Fonte", "Categoria", "Arquivo", "Caminho", "Formato", "Tamanho", "Status técnico", "Licença",
        "Papel da textura", "Resolução", "Modo / canais / alpha", "Power-of-two", "Memória decodificada estimada",
        "Codec", "Duração", "Canais / layout", "Sample rate", "Bitrate", "Família", "Glifos",
        "Membros no contêiner", "Assets internos detectados", "Alertas", "Notas"
    }
    for label, value in compact_rows:
        if label not in keep_labels:
            continue
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.space_after = tools["Pt"](0)
        label_run = paragraph.add_run(f"{label}: ")
        label_run.bold = True
        label_run.font.size = tools["Pt"](6.7)
        value_run = paragraph.add_run(_value_to_text(value))
        value_run.font.size = tools["Pt"](6.7)


def generate_detail_volume(
    records: Sequence[dict[str, Any]],
    output_path: Path,
    title: str,
    subtitle: str,
    cards_per_page: int,
) -> None:
    tools = _import_docx()
    document = tools["Document"]()
    configure_document(document, title, subtitle)
    add_cover(document, title, subtitle, [
        f"Assets neste volume: {len(records):,}",
        f"Origem: {records[0]['origin'] if records else 'N/D'}",
        f"Gerado em: {utc_now_iso()}",
    ])
    if cards_per_page <= 1:
        for record in records:
            add_asset_full_page(document, record)
    else:
        rows_per_page = 2 if cards_per_page == 4 else 1
        columns = 2
        for page_records in _chunks(list(records), cards_per_page):
            table = document.add_table(rows=rows_per_page, cols=columns)
            table.style = "Table Grid"
            table.alignment = tools["WD_TABLE_ALIGNMENT"].CENTER
            for index, record in enumerate(page_records):
                row_index, col_index = divmod(index, columns)
                add_asset_compact_card(table.cell(row_index, col_index), record)
            document.add_page_break()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def generate_master_document(
    records: Sequence[dict[str, Any]],
    summary: dict[str, Any],
    volume_rows: Sequence[dict[str, Any]],
    output_path: Path,
    template_commit: str = "unknown",
    demons_commit: str = "unknown",
) -> None:
    tools = _import_docx()
    document = tools["Document"]()
    title = "Demons Whip - Inventário Mestre de Assets"
    subtitle = "Catálogo técnico e visual do template e dos assets novos"
    configure_document(document, title, subtitle)
    add_cover(document, title, subtitle, [
        f"Assets catalogados: {summary['total_assets']:,}",
        f"Tamanho somado: {human_bytes(summary['total_size_bytes'])}",
        f"Template commit: {template_commit}",
        f"Demons Whip commit: {demons_commit}",
        f"Gerado em: {summary['generated_at']}",
    ])

    document.add_heading("1. Objetivo", level=1)
    document.add_paragraph(
        "Este inventário estabelece uma linha de base técnica e visual para o rebrand Demons Whip. "
        "Ele separa os arquivos que já pertencem ao template World of Claudecraft dos assets novos destinados ao projeto, "
        "mede os dados disponíveis diretamente nos arquivos e prepara a matriz que será usada na etapa seguinte para decidir "
        "o que substitui o quê, o que deve ser otimizado, retexturizado, rerigado, reanimado ou descartado."
    )

    document.add_heading("2. Escopo e metodologia", level=1)
    method_items = [
        "Descoberta por extensão de arquivos de modelos 3D, texturas/imagens, áudio, vídeo, fontes, materiais e contêineres.",
        "Hash SHA-256 para detectar cópias exatas entre packs brutos, assets organizados e template.",
        "Inspeção 3D em Blender para contar objetos, meshes, vértices, faces, triângulos, materiais, UVs, armatures, bones e clips.",
        "Geração de preview padronizado por asset; texturas e imagens preservam transparência sobre fundo quadriculado.",
        "FFprobe/FFmpeg para duração, codec, canais, sample rate, bitrate e previews de áudio/vídeo.",
        "Pillow/fontTools para resolução, canais, alpha, power-of-two, memória estimada e metadados de fontes.",
        "Detecção documental de licença por arquivo próximo e documentos de nível de projeto; aprovação jurídica não é automática.",
    ]
    for item in method_items:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("3. Resumo executivo", level=1)
    add_key_value_table(document, [
        ("Total de assets", f"{summary['total_assets']:,}"),
        ("Tamanho somado", human_bytes(summary["total_size_bytes"])),
        ("Conteúdos únicos por SHA-256", f"{summary['unique_content_hash_count']:,}"),
        ("Cópias em grupos duplicados", f"{summary['duplicate_asset_count']:,}"),
        ("Alertas técnicos", f"{summary['warning_count']:,}"),
        ("Volumes detalhados", f"{len(volume_rows):,}"),
    ])

    document.add_heading("4. Distribuição por origem e tipo", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Origem", "Tipo", "Quantidade", "Participação"]
    for cell, header in zip(table.rows[0].cells, headers):
        _set_cell_shading(cell, "4A171D")
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.color.rgb = tools["RGBColor"].from_string("FFFFFF")
    _set_repeat_table_header(table.rows[0])
    total = max(1, summary["total_assets"])
    for row_data in summary["by_origin_kind"]:
        row = table.add_row()
        values = [row_data["origin"], row_data["kind"], f"{row_data['count']:,}", f"{100 * row_data['count'] / total:.1f}%"]
        for cell, value in zip(row.cells, values):
            cell.text = str(value)

    document.add_heading("5. Status da extração técnica", level=1)
    status_explanations = {
        "OK": "Metadados principais e preview extraídos.",
        "PARTIAL": "Arquivo catalogado, mas parte dos dados depende de inspeção manual ou fonte adicional.",
        "UNSUPPORTED": "Formato sem parser/extrator disponível no pipeline atual.",
        "ERROR": "Arquivo apresentou falha durante leitura, importação ou renderização.",
        "PENDING": "Item não processado; não deveria permanecer após uma execução completa.",
    }
    for status, count in sorted(summary["by_status"].items()):
        document.add_paragraph(f"{status}: {count:,} — {status_explanations.get(status, '')}", style="List Bullet")

    document.add_heading("6. Regras para a próxima etapa", level=1)
    next_rules = [
        "Nenhum asset será considerado substituído sem evidência no runtime.",
        "IDs internos, spawn, collider e footprint do template devem ser preservados quando a mudança for apenas visual.",
        "Polígonos, materiais, texturas, escala, pivô, rig e animações serão comparados lado a lado antes da integração.",
        "Assets sem licença comprovada ficam bloqueados até revisão.",
        "Duplicações serão consolidadas antes de qualquer otimização para evitar trabalho repetido.",
        "Contêineres compilados não serão tratados como fonte editável; será necessário obter os arquivos originais quando forem relevantes.",
    ]
    for item in next_rules:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("7. Índice dos volumes detalhados", level=1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, ["#", "Origem", "Categoria", "Tipo", "Assets", "Arquivo"]):
        _set_cell_shading(cell, "4A171D")
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.color.rgb = tools["RGBColor"].from_string("FFFFFF")
    _set_repeat_table_header(table.rows[0])
    for index, volume in enumerate(volume_rows, start=1):
        row = table.add_row()
        values = [index, volume["origin"], volume["category"], volume["kind"], volume["asset_count"], volume["docx_name"]]
        for cell, value in zip(row.cells, values):
            cell.text = str(value)

    document.add_heading("8. Campos disponíveis por asset", level=1)
    document.add_paragraph(
        "Cada ficha detalhada contém identificação, origem, pack, caminho, formato, tamanho, hash, status técnico, licença, imagem de preview e campos específicos do tipo de asset. "
        "Modelos 3D recebem contagens de geometria, materiais, UVs, rig, bones, animações, dimensões, LODs, colliders e texturas vinculadas. "
        "Texturas recebem resolução, canais, alpha, bit depth estimado, power-of-two e memória descompactada estimada."
    )

    document.add_heading("9. Limitações e ressalvas", level=1)
    limitations = [
        "Root motion, LOD e collider são detectados por dados e convenções de nomes; a validação final ocorre no jogo.",
        "Arquivos .pak, .uasset, .unitypackage e outros contêineres não expõem toda a geometria sem extratores e fontes específicas.",
        "Arquivos .blend podem conter cenas completas com vários assets; a ficha representa o conteúdo do arquivo como unidade técnica.",
        "A memória de textura é uma estimativa sem compressão de GPU; o custo real depende do formato e da plataforma de destino.",
        "A presença de um documento de licença não confirma automaticamente compatibilidade comercial; é apenas evidência para revisão.",
    ]
    for item in limitations:
        document.add_paragraph(item, style="List Bullet")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def generate_all_documents(
    records: list[dict[str, Any]],
    output_dir: Path,
    config: CatalogConfig,
    template_commit: str,
    demons_commit: str,
) -> tuple[Path, list[dict[str, Any]]]:
    volumes_dir = output_dir / "volumes"
    volumes_dir.mkdir(parents=True, exist_ok=True)
    grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        grouped[(record["origin"], record["category"], record["kind"])].append(record)

    volume_rows: list[dict[str, Any]] = []
    for (origin, category, kind), group_records in sorted(grouped.items()):
        group_records.sort(key=lambda item: (item["pack"].lower(), item["relative_path"].lower()))
        if kind == "MODEL_3D":
            max_assets = config.max_model_assets_per_volume
            cards_per_page = 1
        elif kind == "IMAGE_TEXTURE":
            max_assets = config.max_image_assets_per_volume
            cards_per_page = config.image_assets_per_page
        else:
            max_assets = config.max_generic_assets_per_volume
            cards_per_page = config.generic_assets_per_page
        chunks = list(_chunks(group_records, max_assets))
        for volume_index, chunk in enumerate(chunks, start=1):
            stem = f"{origin}_{slugify(category)}_{kind}_V{volume_index:02d}"
            docx_path = volumes_dir / f"Demons_Whip_Catalogo_{stem}.docx"
            title = f"Demons Whip - {origin} - {category}"
            subtitle = f"{kind.replace('_', ' ')} • Volume {volume_index} de {len(chunks)}"
            generate_detail_volume(chunk, docx_path, title, subtitle, cards_per_page)
            volume_rows.append({
                "origin": origin,
                "category": category,
                "kind": kind,
                "volume_index": volume_index,
                "asset_count": len(chunk),
                "docx_path": str(docx_path),
                "docx_name": docx_path.name,
                "pdf_path": str(docx_path.with_suffix(".pdf")),
                "pdf_name": docx_path.with_suffix(".pdf").name,
            })

    summary = summarize_records(records)
    master_path = output_dir / "Demons_Whip_Inventario_Mestre.docx"
    generate_master_document(records, summary, volume_rows, master_path, template_commit, demons_commit)
    return master_path, volume_rows


def convert_docx_files_to_pdf(docx_paths: Sequence[Path], output_parent: Path, timeout_seconds: int = 900) -> list[Path]:
    pdf_paths: list[Path] = []
    for docx_path in docx_paths:
        output_dir = docx_path.parent
        command = [
            "libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(docx_path)
        ]
        env = os.environ.copy()
        env["HOME"] = str(output_parent / ".libreoffice_home")
        Path(env["HOME"]).mkdir(parents=True, exist_ok=True)
        completed = subprocess.run(command, capture_output=True, text=True, timeout=timeout_seconds, check=False, env=env)
        pdf_path = docx_path.with_suffix(".pdf")
        if completed.returncode != 0 or not pdf_path.exists() or pdf_path.stat().st_size == 0:
            raise RuntimeError(
                f"LibreOffice failed for {docx_path}: {completed.stdout}\n{completed.stderr}"
            )
        pdf_paths.append(pdf_path)
    return pdf_paths


def validate_generated_outputs(
    records: Sequence[dict[str, Any]],
    master_docx: Path,
    volume_rows: Sequence[dict[str, Any]],
    data_files: dict[str, Path],
    max_file_mb: int,
) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    expected_asset_ids = {record["asset_id"] for record in records}
    if not master_docx.exists() or master_docx.stat().st_size == 0:
        errors.append("Master DOCX missing or empty")
    inventory_json = json.loads(data_files["json"].read_text(encoding="utf-8"))
    actual_asset_ids = {record["asset_id"] for record in inventory_json}
    missing_ids = sorted(expected_asset_ids - actual_asset_ids)
    extra_ids = sorted(actual_asset_ids - expected_asset_ids)
    if missing_ids:
        errors.append(f"Inventory JSON missing {len(missing_ids)} asset ids")
    if extra_ids:
        errors.append(f"Inventory JSON contains {len(extra_ids)} unexpected asset ids")
    volume_asset_total = sum(int(row["asset_count"]) for row in volume_rows)
    if volume_asset_total != len(records):
        errors.append(f"Volume asset total {volume_asset_total} != inventory total {len(records)}")
    expected_files = [master_docx, master_docx.with_suffix(".pdf")]
    for row in volume_rows:
        expected_files.extend([Path(row["docx_path"]), Path(row["pdf_path"])])
    for path in expected_files:
        if not path.exists() or path.stat().st_size == 0:
            errors.append(f"Missing/empty output: {path}")
            continue
        if path.stat().st_size > max_file_mb * 1024 * 1024:
            warnings.append(f"File exceeds committable threshold: {path.name} ({human_bytes(path.stat().st_size)})")
    preview_missing = [record["asset_id"] for record in records if not Path(record.get("preview_path", "")).exists()]
    if preview_missing:
        errors.append(f"Missing previews for {len(preview_missing)} assets")
    return {
        "ok": not errors,
        "errors": errors,
        "warnings": warnings,
        "asset_count": len(records),
        "volume_asset_total": volume_asset_total,
        "volume_count": len(volume_rows),
        "checked_files": len(expected_files),
    }


def create_delivery_zip(output_dir: Path, zip_path: Path, include_previews: bool = False) -> Path:
    import zipfile
    zip_path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6, allowZip64=True) as archive:
        for path in sorted(output_dir.rglob("*")):
            if not path.is_file():
                continue
            relative = path.relative_to(output_dir)
            if not include_previews and relative.parts and relative.parts[0] == "previews":
                continue
            if any(part.startswith(".") for part in relative.parts):
                continue
            archive.write(path, arcname=normalize_path(Path("Demons_Whip_Asset_Catalog") / relative))
    return zip_path


def write_volume_manifest(volume_rows: Sequence[dict[str, Any]], data_dir: Path) -> Path:
    path = data_dir / "volume_manifest.json"
    cleaned = []
    for row in volume_rows:
        cleaned.append({
            key: (Path(value).name if key.endswith("_path") else value)
            for key, value in row.items()
        })
    path.write_text(json.dumps(cleaned, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def copy_committable_outputs(output_dir: Path, repository_output_dir: Path, max_file_mb: int = 92) -> dict[str, Any]:
    repository_output_dir.mkdir(parents=True, exist_ok=True)
    copied: list[str] = []
    skipped: list[dict[str, Any]] = []
    max_bytes = max_file_mb * 1024 * 1024
    for path in sorted(output_dir.rglob("*")):
        if not path.is_file():
            continue
        relative = path.relative_to(output_dir)
        if relative.parts and relative.parts[0] in {"previews", ".work"}:
            continue
        if path.suffix.lower() == ".zip":
            continue
        if path.stat().st_size > max_bytes:
            skipped.append({"path": normalize_path(relative), "size_bytes": path.stat().st_size})
            continue
        destination = repository_output_dir / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, destination)
        copied.append(normalize_path(relative))
    manifest = {
        "copied": copied,
        "skipped_over_size": skipped,
        "max_file_mb": max_file_mb,
        "generated_at": utc_now_iso(),
    }
    (repository_output_dir / "repository_copy_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return manifest
